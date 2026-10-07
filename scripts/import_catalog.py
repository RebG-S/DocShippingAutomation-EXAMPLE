import os
import re
import sqlite3
import openpyxl

EXCEL_FILE = "Book1.xlsx"
DB_FILE = "catalog.db"


def init_database():
  conn = sqlite3.connect(DB_FILE)
  cursor = conn.cursor()
  cursor.execute("""
    CREATE TABLE IF NOT EXISTS products (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        sku TEXT UNIQUE,
        name TEXT NOT NULL,
        category TEXT DEFAULT 'General',
        unit TEXT DEFAULT 'pcs',
        cost_price REAL DEFAULT 0,
        sell_price REAL NOT NULL,
        reference TEXT,
        updated_at TEXT
    )
    """)
  conn.commit()
  conn.close()


def import_data():
  if not os.path.exists(EXCEL_FILE):
    print(f"[ERROR] File '{EXCEL_FILE}' tidak ditemukan di folder ini!")
    return

  init_database()
  wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
  ws = wb.active

  conn = sqlite3.connect(DB_FILE)
  cursor = conn.cursor()

  imported_count = 0
  print(f"Membaca '{EXCEL_FILE}' sheet '{ws.title}'...")

  for r in range(1, ws.max_row + 1):
    val = ws.cell(r, 1).value
    if not val or not isinstance(val, str) or "|" not in val:
      continue

    parts = [p.strip() for p in val.split("|") if p.strip()]

    # Format baris data: No, Nama Barang, Harga, Tanggal, Ref Sheet
    if len(parts) >= 5:
      no_str, name, price_str, date_str, ref = (
          parts[0],
          parts[1],
          parts[2],
          parts[3],
          parts[4],
      )

      if no_str.lower() in ["no", "#"] or "---" in no_str:
        continue

      # bikin Rp 88.000" -> 88000
      digits = re.sub(r"[^\d]", "", price_str)
      if not digits:
        continue
      clean_price = float(digits)

      # SKU
      words = re.findall(r"[A-Za-z0-9]+", name)
      sku_prefix = (
          "".join([w[:3].upper() for w in words[:3]]) if words else "ITEM"
      )
      sku = f"SKU-{sku_prefix}-{imported_count+1:03d}"

      # Simpan ke SQLite
      cursor.execute(
          """
            INSERT OR REPLACE INTO products (sku, name, sell_price, reference, updated_at)
            VALUES (?, ?, ?, ?, ?)
            """,
          (sku, name, clean_price, ref, date_str),
      )

      imported_count += 1

  conn.commit()
  conn.close()
  print(f"\n[BERHASIL] Total {imported_count} produk sukses masuk ke catalog.db!")


if __name__ == "__main__":
  import_data()