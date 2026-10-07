from datetime import datetime
import io
import json
import os
from dotenv import load_dotenv
import openpyxl
from openpyxl.styles import Alignment
import pandas as pd
import pythoncom
import qrcode
import streamlit as st
from supabase import Client, create_client
import win32com.client

# ==========================================
# 1. KONFIGURASI SUPABASE & KONTAK
# ==========================================
load_dotenv()
SUPABASE_URL = os.getenv("SUPABASE_URL")
SUPABASE_KEY = os.getenv("SUPABASE_KEY")

if not SUPABASE_URL or not SUPABASE_KEY:
  st.error(
      "Kredensial Supabase tidak ditemukan! Pastikan file .env sudah diisi"
      " dengan benar."
  )
  st.stop()

supabase: Client = create_client(SUPABASE_URL, SUPABASE_KEY)

# Metadata Organisasi & Dokumen dari Environment Variables
COMPANY_NAME = os.getenv("COMPANY_NAME", "CV. GLOBAL LOGISTICS")
DEST_DEFAULT = os.getenv("DEST_DEFAULT", "CENTRAL WAREHOUSE")
QUO_PREFIX = os.getenv("QUO_PREFIX", "QUO")
INV_PREFIX = os.getenv("INV_PREFIX", "INV-CORP")
DO_PREFIX = os.getenv("DO_PREFIX", "DO-CORP")

# kontak bersih
DEFAULT_CONTACTS = {
    "Alice": {"name": "Alice Johnson", "email": "alice.j@example.com"},
    "Bob": {"name": "Bob Smith", "email": "bob.s@example.com"},
    "Charlie": {"name": "Charlie Davis", "email": "charlie.d@example.com"},
}

raw_contact_book = os.getenv("CONTACT_BOOK_JSON")
if raw_contact_book:
  try:
    CONTACT_BOOK = json.loads(raw_contact_book)
  except Exception:
    CONTACT_BOOK = DEFAULT_CONTACTS
else:
  CONTACT_BOOK = DEFAULT_CONTACTS

FIRST_CONTACT_KEY = next(iter(CONTACT_BOOK))

st.set_page_config(
    page_title=f"Order Management - {COMPANY_NAME}",
    page_icon="📦",
    layout="wide",
)


# ==========================================
# 2. HELPER SUPABASE DATABASE
# ==========================================
def get_catalog_products():
  try:
    res = (
        supabase.table("products")
        .select("id, sku, name, sell_price, reference, updated_at")
        .order("name")
        .execute()
    )
    return pd.DataFrame(res.data)
  except Exception as e:
    st.error(f"Gagal mengambil katalog dari Supabase: {e}")
    return pd.DataFrame()


def update_or_add_product(name, price, reference="Manual"):
  now_str = datetime.now().strftime("%Y-%m-%d")
  res = (
      supabase.table("products")
      .select("id")
      .ilike("name", name.strip())
      .execute()
  )

  if res.data:
    row_id = res.data[0]["id"]
    supabase.table("products").update({
        "sell_price": price,
        "reference": reference,
        "updated_at": now_str,
    }).eq("id", row_id).execute()
  else:
    sku = f"SKU-{datetime.now().strftime('%y%m%d%H%M%S')}"
    supabase.table("products").insert({
        "sku": sku,
        "name": name.strip(),
        "sell_price": price,
        "reference": reference,
        "updated_at": now_str,
    }).execute()


def load_orders_supabase():
  try:
    res = (
        supabase.table("orders")
        .select("*")
        .order("created_at", desc=True)
        .execute()
    )
    orders_map = {}
    for r in res.data:
      orders_map[r["quo_no"]] = r
    return orders_map
  except Exception as e:
    st.error(f"Gagal memuat daftar order: {e}")
    return {}


def save_order_quo(quo_no, pic_key, items):
  pic_info = CONTACT_BOOK.get(pic_key, CONTACT_BOOK[FIRST_CONTACT_KEY])
  supabase.table("orders").insert({
      "quo_no": quo_no,
      "pic_key": pic_key,
      "pic_name": pic_info["name"],
      "pic_email": pic_info["email"],
      "status": "QUOTED",
      "items": items,
  }).execute()


def update_deal_status(quo_no, inv_no, do_no, po_number):
  supabase.table("orders").update({
      "status": "DEAL",
      "invoice_no": inv_no,
      "do_no": do_no,
      "po_number": po_number,
  }).eq("quo_no", quo_no).execute()


# ==========================================
# 3. HELPER DOKUMEN & WIN32 EXCEL ENGINE
# ==========================================
def get_next_counter(counter_file, start_num):
  if not os.path.exists(counter_file):
    val = start_num
  else:
    with open(counter_file, "r") as f:
      content = f.read().strip()
      val = int(content) if content else start_num
  with open(counter_file, "w") as f:
    f.write(str(val + 1))
  return val


def format_do_date(dt):
  return dt.strftime("%A, %d %B %Y")


def convert_excel_to_pdf(excel_filename):
  pdf_filename = excel_filename.replace(".xlsx", ".pdf")
  abs_excel = os.path.abspath(excel_filename)
  abs_pdf = os.path.abspath(pdf_filename)

  pythoncom.CoInitialize()
  excel = win32com.client.DispatchEx("Excel.Application")
  excel.Visible = False
  excel.DisplayAlerts = False

  try:
    wb = excel.Workbooks.Open(abs_excel)
    wb.ExportAsFixedFormat(0, abs_pdf)
    wb.Close(False)
  finally:
    excel.Quit()
  return pdf_filename


def generate_quotation(pic_key, items):
  wb = openpyxl.load_workbook("QUO-TEMPLATE.xlsx")
  ws = wb.active
  now = datetime.now()
  counter_val = get_next_counter("counter_quo.txt", 101)
  quo_no = f"{QUO_PREFIX}-{now.strftime('%y%m')}-{counter_val}"
  pic_info = CONTACT_BOOK.get(pic_key, CONTACT_BOOK[FIRST_CONTACT_KEY])

  ws["H10"] = now.strftime("%d-%b-%y")
  ws["H11"] = quo_no
  ws["D13"] = pic_info["name"]
  ws["D14"] = pic_info["email"]

  center = Alignment(horizontal="center", vertical="center")
  rows = list(range(19, 50, 2))
  for idx, item in enumerate(items):
    if idx >= len(rows):
      break
    r = rows[idx]
    ws[f"B{r}"] = idx + 1
    ws[f"B{r}"].alignment = center
    ws[f"C{r}"] = item["name"]
    ws[f"F{r}"] = item["price"]
    ws[f"F{r}"].number_format = "#,##0"
    ws[f"G{r}"] = item["qty"]
    ws[f"G{r}"].alignment = center

  excel_file = (
      f"{quo_no}_{pic_info['name'].replace(' ', '_').replace('.', '')}.xlsx"
  )
  wb.save(excel_file)
  save_order_quo(quo_no, pic_key, items)
  pdf_file = convert_excel_to_pdf(excel_file)
  return quo_no, excel_file, pdf_file


def generate_invoice_and_do(quo_no, po_number, do_dt):
  orders_map = load_orders_supabase()
  quo_data = orders_map.get(quo_no)
  now = datetime.now()
  pic_info = CONTACT_BOOK.get(quo_data["pic_key"], CONTACT_BOOK[FIRST_CONTACT_KEY])
  items = quo_data["items"]
  center = Alignment(horizontal="center", vertical="center")

  # --- Invoice ---
  wb_inv = openpyxl.load_workbook("INV-TEMPLATE.xlsx")
  ws_inv = wb_inv.active
  inv_no = f"{INV_PREFIX}-{now.strftime('%y%m')}{get_next_counter('counter_inv.txt', 101)}"
  ws_inv["G12"] = now.strftime("%d-%b-%y")
  ws_inv["G13"] = inv_no
  ws_inv["G14"] = str(po_number)
  ws_inv["C15"] = pic_info["name"]
  ws_inv["C16"] = pic_info["email"]

  for idx, item in enumerate(items):
    r = 21 + (idx * 2)
    ws_inv[f"B{r}"] = idx + 1
    ws_inv[f"B{r}"].alignment = center
    ws_inv[f"C{r}"] = item["name"]
    ws_inv[f"E{r}"] = item["price"]
    ws_inv[f"E{r}"].number_format = "#,##0"
    ws_inv[f"F{r}"] = item["qty"]
    ws_inv[f"F{r}"].alignment = center

  inv_excel = (
      f"{inv_no}_{pic_info['name'].replace(' ', '_').replace('.', '')}.xlsx"
  )
  wb_inv.save(inv_excel)
  inv_pdf = convert_excel_to_pdf(inv_excel)

  # --- Delivery Order ---
  wb_do = openpyxl.load_workbook("DO-Template.xlsx")
  ws_do = wb_do.active
  do_no = f"{DO_PREFIX}-{now.strftime('%y%m')}{get_next_counter('counter_do.txt', 101)}"
  ws_do["C10"] = do_no
  ws_do["F39"] = format_do_date(do_dt)

  for idx, item in enumerate(items):
    r = 15 + (idx * 2)
    ws_do[f"A{r}"] = idx + 1
    ws_do[f"A{r}"].alignment = center
    ws_do[f"B{r}"] = item["name"]
    ws_do[f"G{r}"] = item["qty"]
    ws_do[f"G{r}"].alignment = center

  do_excel = f"{do_no}_OUTBOUND.xlsx"
  wb_do.save(do_excel)
  do_pdf = convert_excel_to_pdf(do_excel)

  update_deal_status(quo_no, inv_no, do_no, po_number)
  return inv_excel, inv_pdf, do_excel, do_pdf


# ==========================================
# HELPER MODAL POP-UP QR CODE UNTUK MOBILE SCANNER
# ==========================================
@st.dialog("🏷️ QR Code Label Pengiriman (80mm Scanner)")
def show_qr_shipping_modal(do_no, po_number, pic_name, items):
  st.write(f"**DO:** `{do_no}` | **PO:** `{po_number}`")
  st.write(f"**Attn:** {pic_name}")

  item_names = [f"{it['name']} (Qty: {it['qty']})" for it in items]
  selected_idx = 0
  if len(items) > 1:
    selected_idx = st.selectbox(
        "Pilih Item untuk Label Dus/Koli:",
        range(len(items)),
        format_func=lambda x: item_names[x],
    )

  chosen_item = items[selected_idx]

  payload = {
      "from": COMPANY_NAME,
      "attn": pic_name,
      "dest": DEST_DEFAULT,
      "po": po_number,
      "qty": chosen_item["qty"],
      "barcode": po_number.replace("/", "").replace("-", "")[:15],
      "desc": chosen_item["name"],
  }

  json_str = json.dumps(payload, separators=(",", ":"))

  qr = qrcode.QRCode(
      version=None,
      error_correction=qrcode.constants.ERROR_CORRECT_M,
      box_size=8,
      border=2,
  )
  qr.add_data(json_str)
  qr.make(fit=True)
  img_qr = qr.make_image(fill_color="black", back_color="white")

  buf = io.BytesIO()
  img_qr.save(buf, format="PNG")
  byte_im = buf.getvalue()

  col_img, col_info = st.columns([1.2, 1])
  with col_img:
    st.image(
        byte_im,
        caption="Arahkan scanner mobile ke QR ini",
        use_container_width=True,
    )
  with col_info:
    st.success("✅ QR Siap Di-scan")
    st.caption("**Isi Payload JSON:**")
    st.code(json_str, language="json")


# ==========================================
# 4. STREAMLIT FRONTEND
# ==========================================
st.title(f"📦 Order & Document Pipeline - {COMPANY_NAME}")

tab1, tab2, tab3 = st.tabs([
    "📝 1. Buat Quotation",
    "🚀 2. Proses Deal (Invoice & DO)",
    "📚 3. Master Katalog & Harga",
])

# ------------------------------------------
# TAB 1: FORM QUOTATION
# ------------------------------------------
with tab1:
  st.subheader("Form Penawaran Baru (Quotation)")

  df_cat = get_catalog_products()
  product_map = {}
  catalog_options = ["-- Pilih Barang dari Katalog --"]

  if not df_cat.empty:
    for _, row in df_cat.iterrows():
      label = row["name"]
      product_map[label] = int(row["sell_price"])
      catalog_options.append(label)

  col_p1, col_p2 = st.columns([1, 2])
  with col_p1:
    pic = st.selectbox("Pilih PIC Rekanan:", list(CONTACT_BOOK.keys()))
  with col_p2:
    st.info(
        f"**Tujuan:** {CONTACT_BOOK[pic]['name']} |"
        f" **Email:** {CONTACT_BOOK[pic]['email']}"
    )

  st.write("---")
  st.write("#### Daftar Item Penawaran")

  num_items = st.number_input(
      "Banyak Item Barang:", min_value=1, max_value=15, value=3
  )

  h_no, h_name, h_price, h_qty = st.columns([0.8, 6.2, 2.5, 1.2])
  with h_no:
    st.markdown("**No**")
  with h_name:
    st.markdown("**Nama Barang / Part Description**")
  with h_price:
    st.markdown("**Harga Satuan (Rp)**")
  with h_qty:
    st.markdown("**Qty**")

  item_inputs = []
  for i in range(int(num_items)):
    c_no, c_name, c_price, c_qty = st.columns([0.8, 6.2, 2.5, 1.2])

    with c_no:
      st.markdown(
          f"<h4 style='text-align: center; margin-top: 5px;'>{i+1}</h4>",
          unsafe_allow_html=True,
      )

    with c_name:
      selected_item = st.selectbox(
          f"Pilih Barang #{i+1}",
          catalog_options,
          key=f"cat_{i}",
          label_visibility="collapsed",
      )

    default_price = (
        product_map[selected_item] if selected_item in product_map else 0
    )

    with c_price:
      item_price = st.number_input(
          f"Harga #{i+1}",
          value=default_price,
          step=5000,
          key=f"price_{i}_{selected_item}",
          label_visibility="collapsed",
      )

    with c_qty:
      item_qty = st.number_input(
          f"Qty #{i+1}",
          min_value=1,
          value=1,
          key=f"qty_{i}",
          label_visibility="collapsed",
      )

    if (
        selected_item != "-- Pilih Barang dari Katalog --"
        and selected_item.strip()
    ):
      item_inputs.append(
          {"name": selected_item.strip(), "price": item_price, "qty": item_qty}
      )

  st.write("---")
  if st.button("📄 Generate Quotation (XLSX + PDF)", type="primary"):
    if not item_inputs:
      st.error("Pilih minimal satu barang dari katalog terlebih dahulu!")
    else:
      with st.spinner("Membuat dokumen Quotation & menyimpan ke Supabase..."):
        quo_no, q_excel, q_pdf = generate_quotation(pic, item_inputs)
        st.success(f"Quotation **{quo_no}** berhasil dibuat dan tersimpan!")

        c_d1, c_d2 = st.columns(2)
        with c_d1:
          with open(q_excel, "rb") as f:
            st.download_button(
                "⬇️ Download Quotation (.xlsx)", f, file_name=q_excel
            )
        with c_d2:
          with open(q_pdf, "rb") as f:
            st.download_button(
                "🖨️ Download Quotation (.pdf) — Siap Cetak",
                f,
                file_name=q_pdf,
                mime="application/pdf",
            )

# ------------------------------------------
# TAB 2: PROSES DEAL (INVOICE & DO)
# ------------------------------------------
with tab2:
  st.subheader("Konfirmasi Deal: Cetak Invoice & Delivery Order")
  orders_map = load_orders_supabase()

  if not orders_map:
    st.info("Belum ada data Quotation di Supabase. Buat dulu di Tab 1.")
  else:
    quo_keys = list(orders_map.keys())
    selected_quo = st.selectbox("Pilih Nomor Quotation:", quo_keys, index=0)
    q_data = orders_map[selected_quo]

    c_m1, c_m2 = st.columns(2)
    with c_m1:
      st.write(f"**PIC:** {q_data.get('pic_name')}")
      st.write(f"**Tanggal Buat:** {q_data.get('created_at')[:10]}")
    with c_m2:
      status_label = (
          "🟢 DEAL" if q_data.get("status") == "DEAL" else "🟡 MENUNGGU PO"
      )
      st.write(f"**Status:** {status_label}")

    with st.expander("👀 Rincian Barang di Penawaran Ini"):
      for idx, it in enumerate(q_data["items"]):
        st.write(
            f"{idx+1}. **{it['name']}** — {it['qty']} pcs @ Rp"
            f" {it['price']:,} = Rp {it['qty']*it['price']:,}"
        )

    st.write("---")
    c_po, c_date = st.columns(2)
    with c_po:
      po_input = st.text_input(
          "Nomor Purchase Order (PO):",
          value=q_data.get("po_number") or "",
          placeholder="Contoh: PO-EXT/2026/09/0142",
      )
    with c_date:
      do_date = st.date_input(
          "Tanggal Delivery Order (DO):", value=datetime.now()
      )

    st.caption(f"Format tanggal cetak DO: **{format_do_date(do_date)}**")

    # EKSEKUSI TERBITKAN INVOICE & DO
    if st.button("⚡ Terbitkan Invoice & DO Sekaligus", type="primary"):
      if not po_input.strip():
        st.error("Nomor PO wajib diisi!")
      else:
        with st.spinner(
            "Membuat file Invoice & DO serta mengupdate Supabase..."
        ):
          inv_xl, inv_pdf, do_xl, do_pdf = generate_invoice_and_do(
              selected_quo, po_input, do_date
          )

          st.session_state[f"deal_{selected_quo}"] = {
              "inv_xl": inv_xl,
              "inv_pdf": inv_pdf,
              "do_xl": do_xl,
              "do_pdf": do_pdf,
              "po": po_input,
              "pic": q_data.get("pic_name"),
              "items": q_data["items"],
          }
          st.rerun()

    deal_key = f"deal_{selected_quo}"
    if deal_key in st.session_state:
      res = st.session_state[deal_key]
      st.success("✅ Dokumen berhasil diterbitkan dan siap diunduh / di-scan!")

      c1, c2 = st.columns(2)
      with c1:
        st.write("📄 **Invoice:**")
        with open(res["inv_xl"], "rb") as f:
          st.download_button("Excel (.xlsx)", f, file_name=res["inv_xl"])
        with open(res["inv_pdf"], "rb") as f:
          st.download_button(
              "🖨️ PDF (Siap Cetak)",
              f,
              file_name=res["inv_pdf"],
              mime="application/pdf",
          )

      with c2:
        st.write("📦 **Delivery Order (DO):**")
        with open(res["do_xl"], "rb") as f:
          st.download_button("Excel (.xlsx)", f, file_name=res["do_xl"])
        with open(res["do_pdf"], "rb") as f:
          st.download_button(
              "🖨️ PDF (Siap Cetak)",
              f,
              file_name=res["do_pdf"],
              mime="application/pdf",
          )

      st.write("---")
      if st.button(
          "📱 Tampilkan QR Label Pengiriman (Scan via HP)",
          type="secondary",
          use_container_width=True,
      ):
        do_code = (
            f"{DO_PREFIX}-{datetime.now().strftime('%y%m')}{get_next_counter('counter_do.txt', 101)-1}"
        )
        show_qr_shipping_modal(
            do_code, res["po"], res["pic"], res["items"]
        )

    elif q_data.get("status") == "DEAL":
      st.write("---")
      if st.button(
          "📱 Buka QR Label Pengiriman untuk Order Ini",
          type="secondary",
          use_container_width=True,
      ):
        show_qr_shipping_modal(
            q_data.get("do_no", f"{DO_PREFIX}-OUT"),
            q_data.get("po_number", "-"),
            q_data.get("pic_name", "-"),
            q_data["items"],
        )

# ------------------------------------------
# TAB 3: MASTER KATALOG & HARGA (SUPABASE)
# ------------------------------------------
with tab3:
  st.subheader("📚 Master Katalog & Histori Harga")

  with st.expander("➕ Update Harga / Tambah Barang Baru"):
    c_u1, c_u2, c_u3 = st.columns([5, 3, 3])
    with c_u1:
      up_name = st.text_input(
          "Nama Barang Lengkap:",
          placeholder="Misal: Heavy Duty Extension Cable 3-Prong",
      )
    with c_u2:
      up_price = st.number_input("Harga Jual (Rp):", min_value=0, step=5000)
    with c_u3:
      up_ref = st.text_input(
          "Referensi Dokumen:", placeholder="Misal: Batch-01 / Ref-102"
      )

    if st.button("💾 Simpan Perubahan ke Supabase"):
      if not up_name.strip() or up_price <= 0:
        st.error("Nama barang dan harga harus diisi!")
      else:
        update_or_add_product(
            up_name.strip(), up_price, up_ref if up_ref.strip() else "Manual"
        )
        st.success(
            f"Data '{up_name}' berhasil disimpan ke cloud Supabase!"
        )
        st.rerun()

  df_display = get_catalog_products()
  search_keyword = st.text_input(
      "🔍 Cari nama barang di katalog:", placeholder="Ketik kata kunci..."
  )

  if not df_display.empty and search_keyword.strip():
    df_display = df_display[
        df_display["name"].str.contains(search_keyword, case=False, na=False)
    ]

  st.caption(f"Menampilkan {len(df_display)} barang dari Supabase")
  if not df_display.empty:
    st.dataframe(
        df_display[["name", "sell_price", "reference", "updated_at"]],
        use_container_width=True,
        column_config={
            "name": st.column_config.TextColumn("Nama Item / Deskripsi"),
            "sell_price": st.column_config.NumberColumn(
                "Harga Satuan", format="Rp %d"
            ),
            "reference": st.column_config.TextColumn("Ref Dokumen"),
            "updated_at": st.column_config.TextColumn("Last Update"),
        },
    )